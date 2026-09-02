"""
VOICE4BLIND — FastAPI Backend
Handles: file listing, PDF/EPUB/DOCX text extraction,
         AI summarization, WebSocket voice pipeline.
"""

import os
import json
import logging
import asyncio
import pathlib
from typing import Optional, List

from fastapi import FastAPI, UploadFile, File, WebSocket, WebSocketDisconnect, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import JSONResponse, FileResponse
from pydantic import BaseModel

# ── Optional heavy deps (graceful fallback) ──────────────────────────────────
try:
    import fitz  # PyMuPDF
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    logging.warning("PyMuPDF not installed — PDF extraction disabled")

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import ebooklib
    from ebooklib import epub
    from bs4 import BeautifulSoup
    EPUB_AVAILABLE = True
except ImportError:
    EPUB_AVAILABLE = False

try:
    import openai
    OPENAI_AVAILABLE = bool(os.environ.get("OPENAI_API_KEY"))
except ImportError:
    OPENAI_AVAILABLE = False

# ── App setup ─────────────────────────────────────────────────────────────────
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("voice4blind")

app = FastAPI(title="VOICE4BLIND API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = pathlib.Path(__file__).parent / "uploads"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

FRONTEND_DIR = pathlib.Path(__file__).parent.parent / "frontend"

# Serve frontend
if FRONTEND_DIR.exists():
    app.mount("/frontend", StaticFiles(directory=str(FRONTEND_DIR), html=True), name="frontend")

# ── User DB (demo) ─────────────────────────────────────────────────────────────
USERS = {
    "harini": "1234",
    "demo":   "demo",
    "user":   "password",
}

# ── File type icons ────────────────────────────────────────────────────────────
ICONS = {".pdf": "📄", ".docx": "📝", ".epub": "📚", ".txt": "📃"}
TYPE_LABELS = {".pdf": "PDF", ".docx": "Word Document", ".epub": "ePub", ".txt": "Text"}

# ─────────────────────────────────────────────────────────────────────────────
# ENDPOINTS
# ─────────────────────────────────────────────────────────────────────────────

@app.get("/")
async def root():
    index = FRONTEND_DIR / "index.html"
    if index.exists():
        return FileResponse(str(index))
    return {"message": "VOICE4BLIND API running. Open /frontend/index.html"}


@app.post("/api/login")
async def login(body: dict):
    username = body.get("username", "").lower().strip()
    password = body.get("password", "").strip()
    stored   = USERS.get(username)
    if stored and stored == password:
        return {"success": True, "username": username}
    raise HTTPException(status_code=401, detail="Invalid credentials")


@app.get("/api/list-files")
async def list_files():
    """List files from uploads directory + scan common user folders."""
    found: List[dict] = []

    # 1. Uploaded files
    for p in sorted(UPLOAD_DIR.iterdir()):
        if p.suffix.lower() in ICONS:
            found.append({
                "name": p.stem.replace("_", " ").title(),
                "path": str(p),
                "type": TYPE_LABELS.get(p.suffix.lower(), p.suffix.upper()),
                "icon": ICONS.get(p.suffix.lower(), "📄"),
            })

    # 2. Common OS folders (non-invasive read-only scan)
    for folder_name in ["Downloads", "Documents", "Desktop"]:
        folder = pathlib.Path.home() / folder_name
        if folder.exists():
            for p in folder.iterdir():
                if p.suffix.lower() in ICONS:
                    found.append({
                        "name": p.stem.replace("_", " ").title(),
                        "path": str(p),
                        "type": TYPE_LABELS.get(p.suffix.lower(), p.suffix.upper()),
                        "icon": ICONS.get(p.suffix.lower(), "📄"),
                    })

    return {"files": found[:20]}  # cap at 20


@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    """Upload a document file to the server."""
    suffix = pathlib.Path(file.filename).suffix.lower()
    if suffix not in ICONS:
        raise HTTPException(400, "Unsupported file type. Please upload PDF, DOCX, EPUB, or TXT.")
    dest = UPLOAD_DIR / file.filename
    content = await file.read()
    dest.write_bytes(content)
    text = extract_text(str(dest))
    return {
        "success": True,
        "filename": file.filename,
        "name": pathlib.Path(file.filename).stem.replace("_", " ").title(),
        "text_preview": text[:300],
    }


@app.get("/api/read-file")
async def read_file(name: str):
    """Extract full text from a named file."""
    # Search uploads + home folders
    candidates = list(UPLOAD_DIR.glob("*"))
    for folder_name in ["Downloads", "Documents", "Desktop"]:
        folder = pathlib.Path.home() / folder_name
        if folder.exists():
            candidates += list(folder.iterdir())

    match = None
    name_lower = name.lower().replace(" ", "")
    for p in candidates:
        if p.suffix.lower() in ICONS:
            if p.stem.lower().replace("_", "").replace(" ", "") == name_lower:
                match = p
                break
            if name_lower in p.stem.lower().replace("_", "").replace(" ", ""):
                match = p

    if not match:
        raise HTTPException(404, f"File not found: {name}")

    text = extract_text(str(match))
    return {"text": text, "filename": match.name}


class SummarizeRequest(BaseModel):
    text: str
    language: Optional[str] = "en"

@app.post("/api/summarize")
async def summarize(req: SummarizeRequest):
    """Summarize text using OpenAI or local fallback."""
    summary = await ai_summarize(req.text, req.language)
    return {"summary": summary}


@app.post("/api/describe-image")
async def describe_image(body: dict):
    """Describe an image/graph from text context."""
    context = body.get("context", "")
    desc = await ai_describe_image(context)
    return {"description": desc}


# ─────────────────────────────────────────────────────────────────────────────
# WEBSOCKET — Real-time voice pipeline
# ─────────────────────────────────────────────────────────────────────────────
@app.websocket("/ws")
async def websocket_endpoint(ws: WebSocket):
    await ws.accept()
    logger.info("WebSocket client connected")
    try:
        while True:
            data = await ws.receive_text()
            msg  = json.loads(data)
            action = msg.get("action")

            if action == "summarize":
                text    = msg.get("text", "")
                lang    = msg.get("language", "en")
                summary = await ai_summarize(text, lang)
                await ws.send_json({"type": "summary", "data": summary})

            elif action == "describe_media":
                context = msg.get("context", "")
                desc    = await ai_describe_image(context)
                await ws.send_json({"type": "media_description", "data": desc})

            elif action == "detect_language":
                text = msg.get("text", "")
                lang = detect_language_hint(text)
                await ws.send_json({"type": "language", "data": lang})

            elif action == "ping":
                await ws.send_json({"type": "pong"})

    except WebSocketDisconnect:
        logger.info("WebSocket client disconnected")


# ─────────────────────────────────────────────────────────────────────────────
# TEXT EXTRACTION
# ─────────────────────────────────────────────────────────────────────────────
def extract_text(path: str) -> str:
    p   = pathlib.Path(path)
    ext = p.suffix.lower()

    if ext == ".pdf":
        return extract_pdf(path)
    elif ext == ".docx":
        return extract_docx(path)
    elif ext == ".epub":
        return extract_epub(path)
    elif ext == ".txt":
        return p.read_text(encoding="utf-8", errors="replace")
    return "Unsupported file format."


def extract_pdf(path: str) -> str:
    if not PDF_AVAILABLE:
        return "PDF extraction not available. Please install PyMuPDF: pip install pymupdf"
    doc  = fitz.open(path)
    text = []
    for page in doc:
        page_text = page.get_text("text")
        text.append(page_text)
        # Detect images / tables on page
        images = page.get_images(full=True)
        if images:
            text.append(f"\n[IMAGE: There are {len(images)} image(s) on this page.]\n")
        tables = page.find_tables()
        if tables and tables.tables:
            for t in tables.tables:
                text.append(f"\n[TABLE: {len(t.rows)} rows × {len(t.cols)} columns]\n")
    return "\n".join(text)


def extract_docx(path: str) -> str:
    if not DOCX_AVAILABLE:
        return "DOCX extraction not available. pip install python-docx"
    doc  = DocxDocument(path)
    text = []
    for para in doc.paragraphs:
        if para.text.strip():
            text.append(para.text)
    for table in doc.tables:
        text.append(f"\n[TABLE: {len(table.rows)} rows × {len(table.columns)} columns]\n")
        for row in table.rows:
            text.append(" | ".join(c.text for c in row.cells))
    return "\n".join(text)


def extract_epub(path: str) -> str:
    if not EPUB_AVAILABLE:
        return "EPUB extraction not available. pip install ebooklib beautifulsoup4"
    book   = epub.read_epub(path)
    chunks = []
    for item in book.get_items():
        if item.get_type() == ebooklib.ITEM_DOCUMENT:
            soup = BeautifulSoup(item.get_content(), "html.parser")
            chunks.append(soup.get_text())
    return "\n".join(chunks)


# ─────────────────────────────────────────────────────────────────────────────
# AI HELPERS
# ─────────────────────────────────────────────────────────────────────────────
async def ai_summarize(text: str, language: str = "en") -> str:
    if OPENAI_AVAILABLE:
        try:
            client = openai.AsyncOpenAI()
            lang_name = {
                "hi": "Hindi", "kn": "Kannada", "ta": "Tamil",
                "te": "Telugu", "ml": "Malayalam", "mr": "Marathi",
            }.get(language[:2], "English")
            resp = await client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[{
                    "role": "system",
                    "content": f"Summarize the following text concisely in {lang_name}. Be brief and clear."
                }, {
                    "role": "user",
                    "content": text[:3000]
                }],
                max_tokens=200,
            )
            return resp.choices[0].message.content
        except Exception as e:
            logger.error(f"OpenAI summarize error: {e}")

    # Fallback: extract first + last sentence
    import re
    sentences = re.findall(r'[^.!?\n]+[.!?\n]+', text)
    if len(sentences) <= 2:
        return text[:300]
    return sentences[0].strip() + " ... " + sentences[-1].strip()


async def ai_describe_image(context: str) -> str:
    if OPENAI_AVAILABLE:
        try:
            client = openai.AsyncOpenAI()
            resp = await client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[{
                    "role": "system",
                    "content": "You are an assistant helping blind students. Describe the chart, graph, or image based on the surrounding document context."
                }, {
                    "role": "user",
                    "content": f"Context: {context[:1000]}\nDescribe what visual element likely appears here."
                }],
                max_tokens=150,
            )
            return resp.choices[0].message.content
        except Exception as e:
            logger.error(f"OpenAI describe error: {e}")
    return "This section contains a visual element such as a chart or diagram. It likely illustrates the data discussed in the surrounding text."


def detect_language_hint(text: str) -> str:
    """Lightweight language detection using character ranges."""
    devanagari = sum(1 for c in text if '\u0900' <= c <= '\u097F')
    kannada    = sum(1 for c in text if '\u0C80' <= c <= '\u0CFF')
    tamil      = sum(1 for c in text if '\u0B80' <= c <= '\u0BFF')
    telugu     = sum(1 for c in text if '\u0C00' <= c <= '\u0C7F')
    total      = len(text)
    if total == 0:
        return "en"
    if kannada / total > 0.2:  return "kn"
    if tamil   / total > 0.2:  return "ta"
    if telugu  / total > 0.2:  return "te"
    if devanagari / total > 0.2: return "hi"
    return "en"


# ─────────────────────────────────────────────────────────────────────────────
# Entry point
# ─────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
